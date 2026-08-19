"""Tests for formulary QA rules and Excel curation pipeline."""
from __future__ import annotations

from pathlib import Path

import pytest

from scripts.formulary.import_xlsx import parse_doses_text
from scripts.formulary.qa_rules import analyze_drug, compute_stable_key


def test_selamectin_hormone_contamination_is_critical():
    drug = {
        "canonical_name_en": "Selamectin",
        "canonical_name_ru": "Selamectin",
        "formulations": "Injectable: 5000 IU freeze-dried plug.",
        "action": "Mimics action of FSH (follicle-stimulating hormone).",
        "use": "Induction of ovulation.",
        "doses": [
            {
                "raw_text": "50–200 IU s.c., i.m. followed by 600 IU hCG after 72 h",
                "source": "bsava",
            }
        ],
    }
    qa = analyze_drug(drug)
    assert qa.critical is True
    assert "antiparasitic_hormone_mix" in qa.flags


def test_pimobendan_antibiotic_contamination_is_critical():
    drug = {
        "id": 1,
        "canonical_name_en": "Pimobendan",
        "canonical_name_ru": "Pimobendan",
        "action": "Beta-lactam antibiotics bind penicillin-binding proteins.",
        "formulations": "2 g piperacillin sodium + 0.25 g tazobactam",
        "doses": [{"raw_text": "0.3 mg/kg PO q12h", "source": "carpenter"}],
    }
    qa = analyze_drug(drug)
    assert qa.critical is True
    assert "inodilator_antibiotic_mix" in qa.flags


def test_selamectin_gets_warning_or_critical_for_iu_doses():
    drug = {
        "id": 2,
        "canonical_name_en": "Selamectin",
        "canonical_name_ru": "Selamectin",
        "action": "Topical antiparasitic.",
        "doses": [{"raw_text": "1000 IU/animal i.m.", "source": "bsava"}],
    }
    qa = analyze_drug(drug)
    assert qa.critical or qa.warning


def test_clean_card_is_not_critical():
    drug = {
        "canonical_name_en": "Meloxicam",
        "canonical_name_ru": "Мелоксикам",
        "action": "NSAID",
        "doses": [{"raw_text": "0.2 mg/kg PO q24h", "source": "manual"}],
    }
    qa = analyze_drug(drug)
    assert qa.critical is False


def test_parse_doses_text_roundtrip():
    text = "[млекопитающие] Guinea pigs: 15 mg/kg topically\n[птицы] Amazon parrot: 10 mg/kg PO q12h"
    doses = parse_doses_text(text)
    assert len(doses) == 2
    assert doses[0]["taxa"] == "mammals"
    assert doses[0]["species_note"] == "Guinea pigs"
    assert "15 mg/kg" in doses[0]["raw_text"]
    assert doses[1]["taxa"] == "birds"


def test_export_import_roundtrip(tmp_path: Path):
    pytest.importorskip("openpyxl")
    from scripts.formulary.export_xlsx import export_workbook
    from scripts.formulary.import_xlsx import import_workbook

    db = Path(__file__).resolve().parents[2] / "data" / "formulary.db"
    if not db.exists():
        pytest.skip("formulary.db not built")

    xlsx = tmp_path / "review.xlsx"
    stats = export_workbook(db, xlsx)
    assert stats["drugs"] > 1000
    assert stats["doses"] > 1000
    assert stats["critical"] >= 1

    from openpyxl import load_workbook

    wb = load_workbook(xlsx)
    ws = wb["Препараты"]
    col_map = {ws.cell(row=2, column=c).value: c for c in range(1, ws.max_column + 1)}
    status_col = col_map["status"]
    comment_col = col_map["comment"]
    critical_col = col_map["qa_critical"]
    for row in range(3, ws.max_row + 1):
        if str(ws.cell(row=row, column=critical_col).value).lower() == "да":
            ws.cell(row=row, column=status_col, value="verified")
            ws.cell(row=row, column=comment_col, value="test import")
            break
    else:
        ws.cell(row=3, column=status_col, value="verified")
        ws.cell(row=3, column=comment_col, value="test import")
    wb.save(xlsx)

    out = tmp_path / "curated.jsonl"
    imported = import_workbook(xlsx, out)
    assert imported["rows"] >= 1
    assert out.exists()


def test_stable_key_matches_inn():
    key = compute_stable_key({"canonical_name_en": "Selamectin"})
    assert key
    assert key == compute_stable_key({"canonical_name_en": "selamectin"})
