"""Export formulary.db → Excel workbook for veterinary curation."""
from __future__ import annotations

import argparse
import json
import logging
import sqlite3
from pathlib import Path
from typing import Any

from scripts.formulary.common import project_root
from scripts.formulary.qa_rules import analyze_all, compute_stable_key

logger = logging.getLogger(__name__)

TAXA_RU = {
    "fish": "рыбы",
    "mammals": "млекопитающие",
    "birds": "птицы",
    "reptiles": "рептилии",
    "amphibians": "амфибии",
    "invertebrates": "беспозвоночные",
    "other": "прочие",
}

DRUG_HEADERS = [
    ("drug_id", "ID (не менять)"),
    ("stable_key", "Ключ (не менять)"),
    ("canonical_name_en", "МНН (EN)"),
    ("canonical_name_ru", "Название (RU)"),
    ("analogs", "Аналоги / торговые"),
    ("formulations", "Формы выпуска"),
    ("action", "Действие"),
    ("use_text", "Применение"),
    ("contraindications", "Противопоказания"),
    ("adverse_reactions", "Побочные эффекты"),
    ("drug_interactions", "Лекарственные взаимодействия"),
    ("safety_handling", "Меры предосторожности"),
    ("pom_note", "POM / статус"),
    ("doses_text", "Дозы (текст)"),
    ("sources", "Источники"),
    ("dose_count", "Кол-во доз (авто)"),
    ("qa_flags", "QA-флаги (авто)"),
    ("qa_critical", "Критично (авто)"),
    ("qa_warning", "Проверить (авто)"),
    ("status", "Статус"),
    ("comment", "Комментарий"),
]

DOSE_HEADERS = [
    ("dose_id", "ID дозы (не менять)"),
    ("stable_key", "Ключ препарата (не менять)"),
    ("drug_id", "ID препарата (не менять)"),
    ("taxa", "Таксон (код)"),
    ("taxa_ru", "Таксон (RU)"),
    ("species_note", "Вид / примечание"),
    ("raw_text", "Текст дозы"),
    ("dose_min", "dose_min"),
    ("dose_max", "dose_max"),
    ("dose_unit", "Единица"),
    ("source", "Источник"),
]

READONLY_DRUG_COLS = {
    "drug_id",
    "stable_key",
    "dose_count",
    "qa_flags",
    "qa_critical",
    "qa_warning",
    "sources",
}

READONLY_DOSE_COLS = {"dose_id", "stable_key", "drug_id", "taxa_ru", "source"}

FILL_CRITICAL = "FFFFC7CE"  # light red
FILL_WARNING = "FFFFEB9C"  # light yellow
FILL_READONLY = "FFF2F2F2"  # light gray
FILL_HEADER = "FFD9E1F2"  # light blue

STATUS_HINT = "todo | in_review | verified | skip"


def _load_drugs(conn: sqlite3.Connection) -> list[dict[str, Any]]:
    conn.row_factory = sqlite3.Row
    drugs: list[dict[str, Any]] = []
    for row in conn.execute(
        """
        SELECT id, canonical_name_en, canonical_name_ru, trade_names, pom_note,
               formulations, action, use_text, safety_handling, contraindications,
               adverse_reactions, drug_interactions, full_text_en, full_text_ru, sources
        FROM drugs
        ORDER BY lower(canonical_name_en), id
        """
    ):
        drug = dict(row)
        drug["trade_names"] = json.loads(drug.get("trade_names") or "[]")
        drug["sources"] = json.loads(drug.get("sources") or "[]")
        drug["use"] = drug.pop("use_text", "") or ""
        doses = [
            dict(d)
            for d in conn.execute(
                """
                SELECT id, taxa, species_note, indication, route,
                       dose_min, dose_max, dose_unit, frequency, duration, raw_text, source
                FROM doses WHERE drug_id=? ORDER BY id
                """,
                (drug["id"],),
            )
        ]
        drug["doses"] = doses
        aliases = [
            r["alias"]
            for r in conn.execute(
                "SELECT alias FROM drug_aliases WHERE drug_id=? ORDER BY alias",
                (drug["id"],),
            )
        ]
        drug["aliases"] = aliases
        drugs.append(drug)
    return drugs


def _format_doses_text(doses: list[dict[str, Any]]) -> str:
    lines: list[str] = []
    for dose in doses:
        taxa = dose.get("taxa") or "other"
        taxa_ru = TAXA_RU.get(taxa, taxa)
        species = (dose.get("species_note") or "").strip()
        raw = (dose.get("raw_text") or "").strip()
        if not raw:
            continue
        prefix = f"[{taxa_ru}]"
        if species:
            lines.append(f"{prefix} {species}: {raw}")
        else:
            lines.append(f"{prefix} {raw}")
    return "\n".join(lines)


def _format_analogs(drug: dict[str, Any]) -> str:
    names = set()
    for value in (
        *(drug.get("trade_names") or []),
        *(drug.get("aliases") or []),
    ):
        cleaned = (value or "").strip()
        if not cleaned:
            continue
        if cleaned in (drug.get("canonical_name_en"), drug.get("canonical_name_ru")):
            continue
        names.add(cleaned)
    return ", ".join(sorted(names, key=str.lower))


def _style_sheet(ws, headers: list[tuple[str, str]], readonly_keys: set[str]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor=FILL_HEADER)
    for col_idx, (key, label) in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        width = 14
        if key in {"doses_text", "action", "use_text", "comment", "qa_flags"}:
            width = 42
        elif key in {"analogs", "formulations", "contraindications"}:
            width = 28
        elif key in {"canonical_name_en", "canonical_name_ru"}:
            width = 22
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    # Store internal keys in hidden row for robust import
    for col_idx, (key, _label) in enumerate(headers, start=1):
        ws.cell(row=2, column=col_idx, value=key)
    ws.row_dimensions[2].hidden = True


def _write_instructions(wb) -> None:
    from openpyxl.styles import Font

    ws = wb.create_sheet("Инструкция", 0)
    lines = [
        "ExoCare+ — курация справочника препаратов",
        "",
        "1. Лист «Препараты» — основная работа.",
        "   • Красные строки = 100% проблемные (склейка PDF, мусор, дубликаты EN).",
        "   • Жёлтые строки = нужна проверка (артефакты PDF, дубликаты имён, подозрительный текст).",
        "2. Начните с фильтра «Критично (авто)» = да, затем «Проверить (авто)» = да.",
        "3. Колонки «(авто)» и «(не менять)» — только для справки, их не редактируют.",
        "4. Заполните «Статус»: todo → in_review → verified (готово) или skip (скрыть из бота).",
        "5. «Комментарий» — для заметок разработчику.",
        "6. Лист «Дозы» — построчная правка доз (опционально). Если правите «Дозы (текст)» на листе Препараты — строки на листе Дозы перезапишутся при импорте.",
        "",
        f"Допустимые статусы: {STATUS_HINT}",
        "",
        "После правок отправьте файл разработчику для импорта и пересборки базы.",
    ]
    for idx, line in enumerate(lines, start=1):
        cell = ws.cell(row=idx, column=1, value=line)
        if idx == 1:
            cell.font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 100


def export_workbook(db_path: Path, out_path: Path) -> dict[str, int]:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill

    if not db_path.exists():
        raise FileNotFoundError(f"Formulary DB not found: {db_path}")

    conn = sqlite3.connect(str(db_path))
    drugs = _load_drugs(conn)
    qa_by_id = analyze_all(drugs)

    wb = Workbook()
    _write_instructions(wb)
    ws_drugs = wb.create_sheet("Препараты")
    wb.remove(wb["Sheet"])

    drug_key_to_row: dict[str, int] = {}
    drug_row_style: dict[int, str] = {}
    critical_fill = PatternFill("solid", fgColor=FILL_CRITICAL)
    warning_fill = PatternFill("solid", fgColor=FILL_WARNING)
    readonly_fill = PatternFill("solid", fgColor=FILL_READONLY)

    _style_sheet(ws_drugs, DRUG_HEADERS, READONLY_DRUG_COLS)
    drug_col_index = {key: idx + 1 for idx, (key, _label) in enumerate(DRUG_HEADERS)}

    for row_idx, drug in enumerate(drugs, start=3):
        stable_key = compute_stable_key(drug)
        drug["stable_key"] = stable_key
        qa = qa_by_id[int(drug["id"])]
        values = {
            "drug_id": drug["id"],
            "stable_key": stable_key,
            "canonical_name_en": drug.get("canonical_name_en") or "",
            "canonical_name_ru": drug.get("canonical_name_ru") or "",
            "analogs": _format_analogs(drug),
            "formulations": drug.get("formulations") or "",
            "action": drug.get("action") or "",
            "use_text": drug.get("use") or "",
            "contraindications": drug.get("contraindications") or "",
            "adverse_reactions": drug.get("adverse_reactions") or "",
            "drug_interactions": drug.get("drug_interactions") or "",
            "safety_handling": drug.get("safety_handling") or "",
            "pom_note": drug.get("pom_note") or "",
            "doses_text": _format_doses_text(drug.get("doses") or []),
            "sources": ", ".join(drug.get("sources") or []),
            "dose_count": len(drug.get("doses") or []),
            "qa_flags": qa.flags_text,
            "qa_critical": "да" if qa.critical else "нет",
            "qa_warning": "да" if qa.warning else "нет",
            "status": "",
            "comment": "",
        }
        for key, col in drug_col_index.items():
            ws_drugs.cell(row=row_idx, column=col, value=values[key])
        if stable_key:
            drug_key_to_row[stable_key] = row_idx
        if qa.critical:
            style = "critical"
            row_fill = critical_fill
        elif qa.warning:
            style = "warning"
            row_fill = warning_fill
        else:
            style = "normal"
            row_fill = None
        drug_row_style[int(drug["id"])] = style
        if row_fill:
            for col in range(1, len(DRUG_HEADERS) + 1):
                ws_drugs.cell(row=row_idx, column=col).fill = row_fill
        else:
            for key in READONLY_DRUG_COLS:
                ws_drugs.cell(row=row_idx, column=drug_col_index[key]).fill = readonly_fill

    ws_doses = wb.create_sheet("Дозы")
    _style_sheet(ws_doses, DOSE_HEADERS, READONLY_DOSE_COLS)
    dose_col_index = {key: idx + 1 for idx, (key, _label) in enumerate(DOSE_HEADERS)}

    dose_row_idx = 3
    dose_count = 0
    for drug in drugs:
        stable_key = drug.get("stable_key") or compute_stable_key(drug)
        for dose in drug.get("doses") or []:
            taxa = dose.get("taxa") or "other"
            values = {
                "dose_id": dose.get("id"),
                "stable_key": stable_key,
                "drug_id": drug["id"],
                "taxa": taxa,
                "taxa_ru": TAXA_RU.get(taxa, taxa),
                "species_note": dose.get("species_note") or "",
                "raw_text": dose.get("raw_text") or "",
                "dose_min": dose.get("dose_min"),
                "dose_max": dose.get("dose_max"),
                "dose_unit": dose.get("dose_unit") or "",
                "source": dose.get("source") or "",
            }
            parent_style = drug_row_style.get(int(drug["id"]), "normal")
            for key, col in dose_col_index.items():
                ws_doses.cell(row=dose_row_idx, column=col, value=values[key])
            if parent_style == "critical":
                row_fill = critical_fill
            elif parent_style == "warning":
                row_fill = warning_fill
            else:
                row_fill = None
            if row_fill:
                for col in range(1, len(DOSE_HEADERS) + 1):
                    ws_doses.cell(row=dose_row_idx, column=col).fill = row_fill
            else:
                for key in READONLY_DOSE_COLS:
                    ws_doses.cell(row=dose_row_idx, column=dose_col_index[key]).fill = readonly_fill
            dose_row_idx += 1
            dose_count += 1

    conn.close()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)

    critical_count = sum(1 for qa in qa_by_id.values() if qa.critical)
    warning_count = sum(1 for qa in qa_by_id.values() if qa.warning)
    stats = {
        "drugs": len(drugs),
        "doses": dose_count,
        "critical": critical_count,
        "warning": warning_count,
    }
    logger.info(
        "Exported %s drugs, %s doses (%s critical, %s warning) → %s",
        stats["drugs"],
        stats["doses"],
        stats["critical"],
        stats["warning"],
        out_path,
    )
    return stats


def main(argv: list[str] | None = None) -> int:
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    root = project_root()
    parser = argparse.ArgumentParser(description="Export formulary.db to Excel for curation")
    parser.add_argument(
        "--db",
        type=Path,
        default=Path(__import__("os").getenv("FORMULARY_DB", str(root / "data" / "formulary.db"))),
    )
    parser.add_argument(
        "--out",
        type=Path,
        default=root / "exports" / "formulary_review.xlsx",
    )
    args = parser.parse_args(argv)
    stats = export_workbook(args.db, args.out)
    print(
        f"Wrote {stats['drugs']} drugs, {stats['doses']} doses "
        f"({stats['critical']} critical / {stats['warning']} warning rows highlighted) "
        f"→ {args.out}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
