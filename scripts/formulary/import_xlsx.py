"""Import curated Excel workbook → formulary_curated.jsonl for DB rebuild."""
from __future__ import annotations

import argparse
import logging
import re
from pathlib import Path
from typing import Any

from scripts.formulary.common import (
    detect_taxa,
    dose_dict,
    empty_drug,
    merge_unique,
    project_root,
    write_jsonl,
)
from scripts.formulary.qa_rules import compute_stable_key

logger = logging.getLogger(__name__)

TAXA_RU_REVERSE = {
    "рыбы": "fish",
    "млекопитающие": "mammals",
    "птицы": "birds",
    "рептилии": "reptiles",
    "амфибии": "amphibians",
    "беспозвоночные": "invertebrates",
    "прочие": "other",
    "fish": "fish",
    "mammals": "mammals",
    "birds": "birds",
    "reptiles": "reptiles",
    "amphibians": "amphibians",
    "invertebrates": "invertebrates",
    "other": "other",
}

DOSE_LINE_RE = re.compile(
    r"^\[(?P<taxa>[^\]]+)\]\s*(?:(?P<species>[^:]+):\s*)?(?P<text>.+)$"
)

VALID_STATUS = frozenset({"todo", "in_review", "verified", "skip", ""})


def _normalize_status(value: Any) -> str:
    text = str(value or "").strip().lower()
    mapping = {
        "todo": "todo",
        "in_review": "in_review",
        "in review": "in_review",
        "verified": "verified",
        "ok": "verified",
        "готово": "verified",
        "skip": "skip",
        "скрыть": "skip",
        "пропустить": "skip",
    }
    return mapping.get(text, text if text in VALID_STATUS else "")


def _sheet_key_map(ws, header_row: int = 2) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        key = ws.cell(row=header_row, column=col).value
        if key:
            mapping[str(key).strip()] = col
    return mapping


def _cell(row: tuple, col_map: dict[str, int], key: str) -> Any:
    idx = col_map.get(key)
    if not idx:
        return None
    return row[idx - 1]


def _parse_analogs(value: Any) -> list[str]:
    text = str(value or "").strip()
    if not text:
        return []
    parts = re.split(r"[,;\n]", text)
    return merge_unique(p.strip() for p in parts if p.strip())


def _parse_taxa(label: str) -> str:
    key = (label or "").strip().lower()
    return TAXA_RU_REVERSE.get(key, detect_taxa(key))


def parse_doses_text(text: str, *, source: str = "curated") -> list[dict[str, Any]]:
    doses: list[dict[str, Any]] = []
    for line in (text or "").splitlines():
        line = line.strip()
        if not line:
            continue
        match = DOSE_LINE_RE.match(line)
        if match:
            taxa = _parse_taxa(match.group("taxa"))
            species = (match.group("species") or "").strip()
            raw = (match.group("text") or "").strip()
        else:
            taxa = detect_taxa(line)
            species = ""
            raw = line
        if not raw:
            continue
        doses.append(
            dose_dict(
                taxa=taxa,
                species_note=species,
                raw_text=raw,
                source=source,
            )
        )
    return doses


def _read_dose_sheet(ws) -> dict[str, list[dict[str, Any]]]:
    col_map = _sheet_key_map(ws)
    if "stable_key" not in col_map:
        raise ValueError("Doses sheet missing stable_key column")
    by_key: dict[str, list[dict[str, Any]]] = {}
    for row_cells in ws.iter_rows(min_row=3, values_only=True):
        if not row_cells or not any(row_cells):
            continue
        stable_key = str(_cell(row_cells, col_map, "stable_key") or "").strip()
        raw_text = str(_cell(row_cells, col_map, "raw_text") or "").strip()
        if not stable_key or not raw_text:
            continue
        taxa = _parse_taxa(str(_cell(row_cells, col_map, "taxa") or "other"))
        species = str(_cell(row_cells, col_map, "species_note") or "").strip()
        dose = dose_dict(
            taxa=taxa,
            species_note=species,
            raw_text=raw_text,
            source="curated",
        )
        dmin = _cell(row_cells, col_map, "dose_min")
        dmax = _cell(row_cells, col_map, "dose_max")
        unit = _cell(row_cells, col_map, "dose_unit")
        if dmin not in (None, ""):
            try:
                dose["dose_min"] = float(dmin)
            except (TypeError, ValueError):
                pass
        if dmax not in (None, ""):
            try:
                dose["dose_max"] = float(dmax)
            except (TypeError, ValueError):
                pass
        if unit:
            dose["dose_unit"] = str(unit).strip()
        by_key.setdefault(stable_key, []).append(dose)
    return by_key


def _read_drug_sheet(ws, doses_by_key: dict[str, list[dict[str, Any]]]) -> list[dict[str, Any]]:
    col_map = _sheet_key_map(ws)
    required = {"stable_key", "canonical_name_en"}
    if not required.issubset(col_map):
        raise ValueError(f"Drugs sheet missing columns: {required - set(col_map)}")

    rows: list[dict[str, Any]] = []
    for row_cells in ws.iter_rows(min_row=3, values_only=True):
        if not row_cells or not any(row_cells):
            continue
        stable_key = str(_cell(row_cells, col_map, "stable_key") or "").strip()
        name_en = str(_cell(row_cells, col_map, "canonical_name_en") or "").strip()
        if not stable_key and not name_en:
            continue
        status = _normalize_status(_cell(row_cells, col_map, "status"))
        comment = str(_cell(row_cells, col_map, "comment") or "").strip()

        # Skip untouched rows: no status and no comment — keep PDF/manual merge as-is.
        if not status and not comment:
            continue

        if not stable_key:
            stable_key = compute_stable_key({"canonical_name_en": name_en})

        analogs = _parse_analogs(_cell(row_cells, col_map, "analogs"))
        name_ru = str(_cell(row_cells, col_map, "canonical_name_ru") or "").strip()
        doses_text = str(_cell(row_cells, col_map, "doses_text") or "").strip()

        doses = doses_by_key.get(stable_key) or []
        if not doses and doses_text:
            doses = parse_doses_text(doses_text, source="curated")

        drug = empty_drug(
            canonical_name_en=name_en or stable_key,
            canonical_name_ru=name_ru,
            source="curated",
        )
        drug["stable_key"] = stable_key
        drug["_stable_key"] = stable_key
        drug["trade_names"] = merge_unique(analogs)
        drug["aliases"] = merge_unique([name_ru, name_en, *analogs])
        drug["formulations"] = str(_cell(row_cells, col_map, "formulations") or "").strip()
        drug["action"] = str(_cell(row_cells, col_map, "action") or "").strip()
        drug["use"] = str(_cell(row_cells, col_map, "use_text") or "").strip()
        drug["contraindications"] = str(_cell(row_cells, col_map, "contraindications") or "").strip()
        drug["adverse_reactions"] = str(_cell(row_cells, col_map, "adverse_reactions") or "").strip()
        drug["drug_interactions"] = str(_cell(row_cells, col_map, "drug_interactions") or "").strip()
        drug["safety_handling"] = str(_cell(row_cells, col_map, "safety_handling") or "").strip()
        drug["pom_note"] = str(_cell(row_cells, col_map, "pom_note") or "").strip()
        drug["doses"] = doses
        drug["status"] = status
        drug["comment"] = comment
        drug["exclude"] = status == "skip"
        drug["sources"] = ["curated"]
        rows.append(drug)
    return rows


def import_workbook(xlsx_path: Path, out_path: Path) -> dict[str, int]:
    from openpyxl import load_workbook

    if not xlsx_path.exists():
        raise FileNotFoundError(f"Excel file not found: {xlsx_path}")

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    if "Препараты" not in wb.sheetnames:
        raise ValueError("Workbook must contain sheet «Препараты»")

    doses_by_key: dict[str, list[dict[str, Any]]] = {}
    if "Дозы" in wb.sheetnames:
        doses_by_key = _read_dose_sheet(wb["Дозы"])

    curated = _read_drug_sheet(wb["Препараты"], doses_by_key)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    count = write_jsonl(out_path, curated)

    stats = {
        "rows": count,
        "verified": sum(1 for r in curated if r.get("status") == "verified"),
        "skip": sum(1 for r in curated if r.get("exclude")),
    }
    logger.info("Imported %s curated rows → %s", count, out_path)
    return stats


def main(argv: list[str] | None = None) -> int:
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    root = project_root()
    parser = argparse.ArgumentParser(description="Import curated Excel → formulary_curated.jsonl")
    parser.add_argument("--xlsx", type=Path, required=True)
    parser.add_argument(
        "--out",
        type=Path,
        default=root / "data" / "formulary_curated.jsonl",
    )
    args = parser.parse_args(argv)
    stats = import_workbook(args.xlsx, args.out)
    print(
        f"Wrote {stats['rows']} curated rows "
        f"({stats['verified']} verified, {stats['skip']} skip) → {args.out}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
